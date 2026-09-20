from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import timedelta
from pathlib import Path
from typing import Iterable

from django.conf import settings
from django.db.models import QuerySet
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics import renderPDF
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from svglib.svglib import svg2rlg

from website.models import AuditEvent, ExportRecord, ServiceRequest, SiteSettings


CW_BLUE = colors.HexColor("#0754A6")
DEEP_BLUE = colors.HexColor("#063B78")
CW_ORANGE = colors.HexColor("#F5A000")
CW_GOLD = colors.HexColor("#FDBA18")
BACKGROUND = colors.HexColor("#F7F9FC")
TEXT = colors.HexColor("#152033")

EXPORT_COLUMNS = (
    ("ticket_number", "Ticket"),
    ("created_at", "Fecha"),
    ("name", "Cliente"),
    ("whatsapp", "WhatsApp"),
    ("email", "Email"),
    ("service", "Servicio"),
    ("equipment", "Equipo"),
    ("brand", "Marca"),
    ("model", "Modelo"),
    ("municipality", "Municipio"),
    ("sector", "Zona"),
    ("status", "Estado"),
    ("assigned_technician", "Técnico"),
    ("source", "Origen"),
)


def csv_safe(value) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _display_value(item: ServiceRequest, field: str):
    if field == "created_at":
        return timezone.localtime(item.created_at).replace(tzinfo=None)
    if field == "status":
        return item.get_status_display()
    if field == "service":
        return item.service.name if item.service else ""
    if field == "assigned_technician":
        return item.assigned_technician.get_full_name() or item.assigned_technician.get_username() if item.assigned_technician else ""
    return getattr(item, field, "")


def service_request_rows(queryset: QuerySet[ServiceRequest]) -> list[list]:
    return [
        [_display_value(item, field) for field, _label in EXPORT_COLUMNS]
        for item in queryset.select_related("service", "assigned_technician").iterator(chunk_size=500)
    ]


def build_csv(rows: Iterable[Iterable]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow([label for _field, label in EXPORT_COLUMNS])
    for row in rows:
        writer.writerow([csv_safe(value.isoformat(sep=" ") if hasattr(value, "isoformat") else value) for value in row])
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def build_json(rows: Iterable[Iterable]) -> bytes:
    labels = [label for _field, label in EXPORT_COLUMNS]
    records = []
    for row in rows:
        records.append(
            {
                label: value.isoformat() if hasattr(value, "isoformat") else value
                for label, value in zip(labels, row, strict=True)
            }
        )
    return json.dumps(records, ensure_ascii=False, indent=2).encode("utf-8")


def build_xlsx(rows: list[list], generated_by: str) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary.sheet_view.showGridLines = False
    summary["A1"] = "CW REPARACIONES"
    summary["A1"].font = Font(size=20, bold=True, color="0754A6")
    summary["A2"] = "Reporte de solicitudes"
    summary["A2"].font = Font(size=14, bold=True, color="152033")
    summary["A4"] = "Registros"
    summary["B4"] = len(rows)
    summary["A5"] = "Generado el"
    summary["B5"] = timezone.localtime().replace(tzinfo=None)
    summary["B5"].number_format = "yyyy-mm-dd hh:mm"
    summary["A6"] = "Generado por"
    summary["B6"] = generated_by
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 32

    sheet = workbook.create_sheet("Solicitudes")
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    headers = [label for _field, label in EXPORT_COLUMNS]
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="063B78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24
    widths = (20, 19, 28, 18, 30, 26, 24, 18, 18, 18, 18, 22, 24, 20)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet["B"][1:]:
        cell.number_format = "yyyy-mm-dd hh:mm"
    for row_index in range(2, sheet.max_row + 1):
        if row_index % 2 == 0:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F7F9FC")
        for cell in sheet[row_index]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        page_count = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.setFillColor(colors.HexColor("#5C667A"))
            self.setFont("Helvetica", 8)
            self.drawRightString(self._pagesize[0] - 14 * mm, 10 * mm, f"Página {self._pageNumber} de {page_count}")
            super().showPage()
        super().save()


class CWBrandedPDFTemplate:
    """Plantilla corporativa única para todos los PDF oficiales de CW."""

    def __init__(self, *, title: str, subject: str, internal: bool = False, landscape_mode: bool = False):
        self.title = title
        self.subject = subject
        self.internal = internal
        self.pagesize = landscape(A4) if landscape_mode else A4
        self.styles = getSampleStyleSheet()
        self.styles.add(ParagraphStyle(name="CWTitle", parent=self.styles["Title"], textColor=DEEP_BLUE, fontSize=19, leading=23, spaceAfter=10))
        self.styles.add(ParagraphStyle(name="CWSection", parent=self.styles["Heading2"], textColor=DEEP_BLUE, fontSize=12, leading=15, spaceBefore=8, spaceAfter=7))
        self.styles.add(ParagraphStyle(name="CWKPI", parent=self.styles["BodyText"], alignment=TA_CENTER, textColor=DEEP_BLUE, fontSize=15, leading=18))

    def _logo(self):
        logo_path = Path(settings.BASE_DIR) / "static" / "website" / "images" / "cw-logo.svg"
        drawing = svg2rlg(str(logo_path))
        if drawing is None:
            return None
        max_width, max_height = 47 * mm, 16 * mm
        factor = min(max_width / drawing.width, max_height / drawing.height)
        drawing.scale(factor, factor)
        drawing.width *= factor
        drawing.height *= factor
        return drawing

    def _on_page(self, canv, doc) -> None:
        width, height = self.pagesize
        canv.saveState()
        canv.setTitle(self.title)
        canv.setAuthor("CW Reparaciones")
        canv.setSubject(self.subject)
        canv.setFillColor(DEEP_BLUE)
        canv.rect(0, height - 28 * mm, width, 28 * mm, stroke=0, fill=1)
        logo = self._logo()
        if logo:
            canv.setFillColor(colors.white)
            canv.roundRect(10 * mm, height - 25 * mm, 54 * mm, 20 * mm, 3 * mm, stroke=0, fill=1)
            renderPDF.draw(logo, canv, 14 * mm, height - 22 * mm)
        canv.setFillColor(colors.white)
        canv.setFont("Helvetica-Bold", 12)
        canv.drawRightString(width - 14 * mm, height - 13 * mm, "CW REPARACIONES")
        canv.setFont("Helvetica", 8)
        canv.drawRightString(width - 14 * mm, height - 18 * mm, "Servicios Técnicos")
        canv.setFillColor(CW_ORANGE)
        canv.rect(0, height - 29 * mm, width, 1.2 * mm, stroke=0, fill=1)
        if self.internal:
            canv.setFillColor(CW_ORANGE)
            canv.setFont("Helvetica-Bold", 8)
            canv.drawString(14 * mm, 10 * mm, "USO INTERNO")
        canv.setFillColor(colors.HexColor("#5C667A"))
        canv.setFont("Helvetica", 8)
        canv.drawCentredString(width / 2, 10 * mm, "CW Reparaciones - Servicios Técnicos - +57 317 504 0053 - @cwreparaciones")
        canv.restoreState()

    def build(self, story: list) -> bytes:
        output = io.BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=self.pagesize,
            rightMargin=14 * mm,
            leftMargin=14 * mm,
            topMargin=38 * mm,
            bottomMargin=18 * mm,
            title=self.title,
            author="CW Reparaciones",
            subject=self.subject,
        )
        document.build(story, onFirstPage=self._on_page, onLaterPages=self._on_page, canvasmaker=NumberedCanvas)
        return output.getvalue()


def _pdf_table(data: list[list], widths=None) -> Table:
    table = Table(data, repeatRows=1, colWidths=widths)
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), DEEP_BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9E1EC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for index in range(2, len(data), 2):
        commands.append(("BACKGROUND", (0, index), (-1, index), BACKGROUND))
    table.setStyle(TableStyle(commands))
    return table


def build_report_pdf(rows: list[list], generated_by: str) -> bytes:
    template = CWBrandedPDFTemplate(
        title="Reporte de solicitudes",
        subject="Reporte operativo de solicitudes de servicio",
        internal=True,
        landscape_mode=True,
    )
    total = len(rows)
    completed = sum(1 for row in rows if row[11] == "Finalizada")
    scheduled = sum(1 for row in rows if row[11] == "Programada")
    pending = total - completed
    story = [
        Paragraph("REPORTE DE SOLICITUDES", template.styles["CWTitle"]),
        Paragraph(
            f"Generado el {timezone.localtime():%d/%m/%Y %H:%M} - Generado por {generated_by}",
            template.styles["BodyText"],
        ),
        Spacer(1, 6 * mm),
        Table(
            [
                [Paragraph(f"<b>{total}</b><br/>Solicitudes", template.styles["CWKPI"]), Paragraph(f"<b>{completed}</b><br/>Finalizadas", template.styles["CWKPI"]), Paragraph(f"<b>{pending}</b><br/>Pendientes", template.styles["CWKPI"]), Paragraph(f"<b>{scheduled}</b><br/>Programadas", template.styles["CWKPI"])],
            ],
            colWidths=[64 * mm] * 4,
            style=TableStyle([("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#D9E1EC")), ("BACKGROUND", (0, 0), (-1, -1), BACKGROUND), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]),
        ),
        Spacer(1, 7 * mm),
        Paragraph("Detalle autorizado", template.styles["CWSection"]),
    ]
    selected_indexes = (0, 1, 2, 5, 6, 7, 9, 11, 12)
    headers = [EXPORT_COLUMNS[index][1] for index in selected_indexes]
    pdf_rows = [[str(row[index]) if row[index] is not None else "" for index in selected_indexes] for row in rows]
    story.append(_pdf_table([headers, *pdf_rows]))
    return template.build(story)


def build_service_order_pdf(service_request: ServiceRequest, *, internal: bool) -> bytes:
    title = f"Orden de servicio {service_request.ticket_number}"
    template = CWBrandedPDFTemplate(title=title, subject="Orden de servicio", internal=internal)
    contact = service_request.whatsapp or service_request.phone
    details = [
        ["Ticket", service_request.ticket_number, "Estado", service_request.get_status_display()],
        ["Fecha", timezone.localtime(service_request.created_at).strftime("%d/%m/%Y %H:%M"), "Preferencia", service_request.get_contact_preference_display()],
        ["Cliente", service_request.name, "Teléfono", contact],
        ["Zona", f"{service_request.municipality} · {service_request.sector}".strip(" ·"), "Visita", timezone.localtime(service_request.scheduled_start).strftime("%d/%m/%Y %H:%M") if service_request.scheduled_start else "No programada"],
        ["Equipo", service_request.equipment, "Marca / modelo", f"{service_request.brand} {service_request.model}".strip()],
    ]
    story = [
        Paragraph("ORDEN DE SERVICIO", template.styles["CWTitle"]),
        _pdf_table([["Campo", "Información", "Campo", "Información"], *details], widths=[30 * mm, 55 * mm, 30 * mm, 55 * mm]),
        Spacer(1, 7 * mm),
        Paragraph("Problema reportado", template.styles["CWSection"]),
        Paragraph(service_request.description or service_request.issue, template.styles["BodyText"]),
    ]
    if internal:
        story.extend([
            Paragraph("Información operativa", template.styles["CWSection"]),
            Paragraph(service_request.technical_notes or "Sin notas técnicas.", template.styles["BodyText"]),
            Paragraph(service_request.internal_notes or "Sin notas internas.", template.styles["BodyText"]),
        ])
    return template.build(story)


def create_service_request_export(*, user, queryset: QuerySet[ServiceRequest], file_format: str, parameters: dict | None = None) -> ExportRecord:
    if file_format not in ExportRecord.Format.values:
        raise ValueError("Formato de exportación no permitido.")
    expires_at = timezone.now() + timedelta(hours=settings.WEBSITE_EXPORT_RETENTION_HOURS)
    record = ExportRecord.objects.create(
        created_by=user,
        file_format=file_format,
        parameters=parameters or {},
        expires_at=expires_at,
    )
    try:
        rows = service_request_rows(queryset)
        username = user.get_full_name() or user.get_username()
        builders = {
            ExportRecord.Format.CSV: lambda: build_csv(rows),
            ExportRecord.Format.XLSX: lambda: build_xlsx(rows, username),
            ExportRecord.Format.PDF: lambda: build_report_pdf(rows, username),
            ExportRecord.Format.JSON: lambda: build_json(rows),
        }
        content = builders[file_format]()
        stamp = timezone.localtime().strftime("%Y%m%d-%H%M%S")
        filename = f"cw-reparaciones-solicitudes-{stamp}.{file_format}"
        relative = Path(str(user.pk)) / timezone.localdate().isoformat() / str(record.pk) / filename
        root = Path(settings.PRIVATE_EXPORT_ROOT).resolve()
        target = (root / relative).resolve()
        if root not in target.parents:
            raise ValueError("Ruta de exportación no segura.")
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_suffix(target.suffix + ".tmp")
        temporary.write_bytes(content)
        temporary.replace(target)
        record.status = ExportRecord.Status.READY
        record.relative_path = relative.as_posix()
        record.original_filename = filename
        record.row_count = len(rows)
        record.size_bytes = len(content)
        record.sha256 = hashlib.sha256(content).hexdigest()
        record.save(update_fields=("status", "relative_path", "original_filename", "row_count", "size_bytes", "sha256"))
        AuditEvent.objects.create(
            actor=user,
            action=AuditEvent.Action.EXPORT,
            object_type="website.exportrecord",
            object_id=str(record.pk),
            object_repr=filename,
            changes={"format": file_format, "rows": len(rows), "parameters": parameters or {}},
        )
        return record
    except Exception as exc:
        record.status = ExportRecord.Status.FAILED
        record.error_message = str(exc)[:500]
        record.save(update_fields=("status", "error_message"))
        raise


def resolve_export_path(record: ExportRecord) -> Path:
    root = Path(settings.PRIVATE_EXPORT_ROOT).resolve()
    target = (root / record.relative_path).resolve()
    if not record.relative_path or root not in target.parents:
        raise ValueError("Ruta de exportación no segura.")
    return target
