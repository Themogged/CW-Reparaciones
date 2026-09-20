from __future__ import annotations

import json
import re
from pathlib import Path

from django.contrib.auth import get_user_model
from django.conf import settings
from django.core.exceptions import ValidationError
from django.test import TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from website.business.exports import (
    build_service_order_pdf,
    create_service_request_export,
    csv_safe,
)
from website.business.backups import create_backup
from website.models import (
    AuditEvent,
    BackupRecord,
    Customer,
    ExportRecord,
    ServiceRequest,
    ServiceRequestEvent,
    UserProfile,
)


def make_request(**overrides) -> ServiceRequest:
    values = {
        "equipment": "Lavadora",
        "issue": "No centrifuga",
        "frequency": ServiceRequest.Frequency.CONSTANT,
        "description": "La lavadora no completa el ciclo.",
        "name": "Cliente real de prueba",
        "whatsapp": "+57 300 123 4567",
        "contact_preference": ServiceRequest.ContactPreference.WHATSAPP,
        "privacy_accepted": True,
    }
    values.update(overrides)
    return ServiceRequest.objects.create(**values)


class TicketAndCustomerTests(TestCase):
    def test_new_requests_receive_sequential_yearly_ticket(self):
        first = make_request(whatsapp="+57 300 000 0001")
        second = make_request(whatsapp="+57 300 000 0002")
        year = timezone.localdate().year
        self.assertEqual(first.ticket_number, f"CW-{year}-000001")
        self.assertEqual(second.ticket_number, f"CW-{year}-000002")
        self.assertRegex(second.ticket_number, re.compile(r"^CW-\d{4}-\d{6}$"))

    def test_exact_contact_reuses_customer_profile(self):
        first = make_request(whatsapp="+57 300 999 8877")
        second = make_request(whatsapp="3009998877")
        self.assertEqual(first.customer_id, second.customer_id)
        self.assertEqual(Customer.objects.count(), 1)

    def test_status_change_creates_timeline_event(self):
        item = make_request()
        item.status = ServiceRequest.Status.IN_PROGRESS
        item.save()
        self.assertTrue(item.timeline.filter(event_type="status_changed").exists())


class RoleSecurityTests(TestCase):
    def setUp(self):
        self.User = get_user_model()

    def _user_with_role(self, username, role):
        user = self.User.objects.create_user(username=username, password="A-strong-passphrase-2026", is_staff=True)
        profile = user.cw_profile
        profile.role = role
        profile.save()
        user.refresh_from_db()
        return user

    def test_unassigned_new_staff_has_no_private_request_permission(self):
        user = self.User.objects.create_user(username="pending", password="A-strong-passphrase-2026", is_staff=True)
        self.assertFalse(user.has_perm("website.view_servicerequest"))

    def test_technician_cannot_download_unassigned_attachment(self):
        technician = self._user_with_role("tech", UserProfile.Role.TECHNICIAN)
        item = make_request()
        self.client.force_login(technician)
        response = self.client.get(reverse("website:request_attachment_download", args=(item.pk,)))
        self.assertEqual(response.status_code, 403)

    def test_last_owner_role_is_protected(self):
        owner = self._user_with_role("owner", UserProfile.Role.OWNER)
        profile = owner.cw_profile
        profile.role = UserProfile.Role.ADMIN
        with self.assertRaises(ValidationError):
            profile.save()


class ExportTests(TestCase):
    def setUp(self):
        self.export_root = Path(settings.BASE_DIR) / "private_exports" / "automated-tests"
        self.export_root.mkdir(parents=True, exist_ok=True)
        self.override = override_settings(PRIVATE_EXPORT_ROOT=self.export_root)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.user = get_user_model().objects.create_user(
            username="exporter",
            password="A-strong-passphrase-2026",
            is_staff=True,
        )
        profile = self.user.cw_profile
        profile.role = UserProfile.Role.OWNER
        profile.save()
        self.item = make_request(name="=CMD|cliente", whatsapp="+57 300 111 2233")

    def test_csv_formula_injection_is_neutralized(self):
        self.assertEqual(csv_safe("=1+1"), "'=1+1")
        record = create_service_request_export(
            user=self.user,
            queryset=ServiceRequest.objects.all(),
            file_format=ExportRecord.Format.CSV,
        )
        exported = (self.export_root / record.relative_path).read_text(encoding="utf-8-sig")
        self.assertIn("'=CMD|cliente", exported)

    def test_xlsx_has_professional_structure(self):
        record = create_service_request_export(
            user=self.user,
            queryset=ServiceRequest.objects.all(),
            file_format=ExportRecord.Format.XLSX,
        )
        workbook = load_workbook(self.export_root / record.relative_path)
        self.assertEqual(workbook.sheetnames, ["Resumen", "Solicitudes"])
        self.assertEqual(workbook["Solicitudes"].freeze_panes, "A2")
        self.assertTrue(workbook["Solicitudes"].auto_filter.ref)

    def test_pdf_uses_real_logo_and_is_generated(self):
        content = build_service_order_pdf(self.item, internal=False)
        self.assertTrue(content.startswith(b"%PDF"))
        self.assertGreater(len(content), 3_000)

    def test_json_is_valid_utf8(self):
        record = create_service_request_export(
            user=self.user,
            queryset=ServiceRequest.objects.all(),
            file_format=ExportRecord.Format.JSON,
        )
        data = json.loads((self.export_root / record.relative_path).read_text(encoding="utf-8"))
        self.assertEqual(data[0]["Ticket"], self.item.ticket_number)

    def test_secure_download_checks_owner_and_integrity(self):
        record = create_service_request_export(
            user=self.user,
            queryset=ServiceRequest.objects.all(),
            file_format=ExportRecord.Format.CSV,
        )
        stranger = get_user_model().objects.create_user(username="stranger", password="A-strong-passphrase-2026", is_staff=True)
        self.client.force_login(stranger)
        self.assertEqual(self.client.get(reverse("website:export_download", args=(record.token,))).status_code, 403)
        self.client.force_login(self.user)
        response = self.client.get(reverse("website:export_download", args=(record.token,)))
        self.assertEqual(response.status_code, 200)
        response.close()


class AuditImmutabilityTests(TestCase):
    def test_audit_event_cannot_be_updated_or_deleted(self):
        event = AuditEvent.objects.create(action=AuditEvent.Action.CREATE, object_type="test")
        event.object_repr = "changed"
        with self.assertRaises(ValidationError):
            event.save()
        with self.assertRaises(ValidationError):
            event.delete()


class AdminDashboardTests(TestCase):
    def test_dashboard_uses_real_counts_and_renders(self):
        owner = get_user_model().objects.create_superuser(
            username="dashboard-owner",
            password="A-strong-passphrase-2026",
            email="owner@example.test",
        )
        make_request()
        self.client.force_login(owner)
        response = self.client.get(reverse("admin:index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Resumen del negocio")
        self.assertContains(response, "CW-")


class ManualServiceRequestAdminTests(TestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_superuser(
            username="manual-owner",
            password="A-strong-passphrase-2026",
            email="owner@example.test",
        )
        self.client.force_login(self.owner)

    def _post_data(self, **overrides):
        values = {
            "version_token": "1",
            "manual_channel": "phone",
            "status": ServiceRequest.Status.NEW,
            "equipment": "Nevera",
            "brand": "Marca comprobada",
            "model": "Modelo 100",
            "issue": "No enfría",
            "frequency": ServiceRequest.Frequency.CONSTANT,
            "description": "El cliente reportó que el equipo dejó de enfriar.",
            "name": "Cliente telefónico",
            "phone": "300 555 0101",
            "whatsapp": "300 555 0101",
            "email": "",
            "municipality": "Medellín",
            "sector": "Centro",
            "address": "",
            "preferred_date": "",
            "preferred_time": "",
            "contact_preference": ServiceRequest.ContactPreference.PHONE,
            "privacy_accepted": "on",
            "technical_notes": "",
            "internal_notes": "",
            "notes-TOTAL_FORMS": "0",
            "notes-INITIAL_FORMS": "0",
            "notes-MIN_NUM_FORMS": "0",
            "notes-MAX_NUM_FORMS": "1000",
            "timeline-TOTAL_FORMS": "0",
            "timeline-INITIAL_FORMS": "0",
            "timeline-MIN_NUM_FORMS": "0",
            "timeline-MAX_NUM_FORMS": "0",
            "_save": "Guardar",
        }
        values.update(overrides)
        return values

    def test_changelist_exposes_clear_manual_creation_action(self):
        response = self.client.get(reverse("admin:website_servicerequest_changelist"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Nueva solicitud manual")

    def test_owner_can_create_manual_request_and_open_its_orders(self):
        response = self.client.post(
            reverse("admin:website_servicerequest_add"),
            self._post_data(),
        )
        item = ServiceRequest.objects.get()
        self.assertRedirects(
            response,
            reverse("admin:website_servicerequest_change", args=(item.pk,)),
        )
        self.assertEqual(item.source, "manual:phone")
        self.assertTrue(item.ticket_number.startswith("CW-"))
        self.assertIsNotNone(item.customer_id)
        self.assertTrue(item.privacy_accepted)
        self.assertIsNotNone(item.consented_at)
        created_event = item.timeline.get(event_type=ServiceRequestEvent.EventType.CREATED)
        self.assertEqual(created_event.actor, self.owner)
        self.assertIn("llamada telefónica", created_event.description.lower())
        self.assertTrue(
            AuditEvent.objects.filter(
                actor=self.owner,
                action=AuditEvent.Action.CREATE,
                object_type="website.servicerequest",
                object_id=str(item.pk),
            ).exists()
        )

    def test_manual_request_requires_customer_authorization(self):
        response = self.client.post(
            reverse("admin:website_servicerequest_add"),
            self._post_data(privacy_accepted=""),
        )
        self.assertEqual(
            response.status_code,
            200,
            response.content.decode("utf-8", errors="replace")[:2000],
        )
        self.assertContains(response, "Este campo es obligatorio")
        self.assertFalse(ServiceRequest.objects.exists())


class BackupTests(TransactionTestCase):
    def test_database_backup_has_integrity_hash(self):
        backup_root = Path(settings.BASE_DIR) / "private_backups" / "automated-tests"
        backup_root.mkdir(parents=True, exist_ok=True)
        user = get_user_model().objects.create_superuser(
            username="backup-owner",
            password="A-strong-passphrase-2026",
            email="backup@example.test",
        )
        with override_settings(PRIVATE_BACKUP_ROOT=backup_root):
            record = create_backup(user=user, scope=BackupRecord.Scope.DATABASE)
        path = backup_root / record.relative_path
        self.assertEqual(record.status, BackupRecord.Status.READY)
        self.assertTrue(path.is_file())
        self.assertEqual(len(record.sha256), 64)
